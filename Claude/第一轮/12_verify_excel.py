"""步骤10：复核最终 Excel —— 工作表齐全、能读出、关键数字与汇总表一致、内部加总自洽。"""
import openpyxl, pandas as pd, pathlib

BASE = pathlib.Path(r"D:\项目2\Claude")
AGG = BASE / "中间汇总"
wb = openpyxl.load_workbook(BASE / "最终成果.xlsx")
def L(n): return pd.read_csv(AGG / f"{n}.csv", encoding="utf-8-sig")

print("工作表数:", len(wb.sheetnames))
for s in wb.sheetnames:
    w = wb[s]
    print(f"  {s:<18s} 行={w.max_row:<4d} 列={w.max_column:<3d} 图表={len(getattr(w,'_charts',[]))}")

checks = []
def chk(name, got, exp):
    checks.append((name, str(got), str(exp)))

# ---- Excel 单元格 vs CSV ----
ws = wb["2 数据概况"]
cells = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, 30)}
m01 = dict(zip(L("t01_overview")["指标"], L("t01_overview")["值"].astype(str)))
for k in ["数据集原始记录数", "本看板分析口径记录数", "剔除的异常聚集记录数（见 3.2）", "覆盖天数"]:
    chk(f"概况/{k}", cells.get(k), m01[k])

ws = wb["4 月度趋势"]
found = {}
for row in ws.iter_rows(values_only=True):
    if row and isinstance(row[0], str) and row[0].startswith(("Y1 ", "Y2 ")):
        found[row[0][:2]] = (row[1], row[3])
t05 = L("t05_yoy_window")
for _, rw in t05.iterrows():
    k = rw["窗口"][:2]
    chk(f"同比/{k} 记录数", found[k][0], rw["记录数"])
    chk(f"同比/{k} 日均", f"{float(found[k][1]):.1f}", f"{float(rw['日均记录数']):.1f}")

# ---- 口径自洽 ----
N_ALL, N_EXCL, N_ANA = 7525498, 185791, 7339707
chk("原始 - 剔除 = 分析口径", N_ALL - N_EXCL, N_ANA)
chk("剔除量 = t28 合计", int(L("t28_excluded_pairs")["两年合计记录数"].sum()), N_EXCL)

t04 = L("t04_monthly")
chk("月度记录数合计 = 分析口径", int(t04["记录数"].sum()), N_ANA)
chk("完整月个数", (t04["月份完整性"] == "完整月").sum(), 23)

n1, n2 = int(t05.loc[0, "记录数"]), int(t05.loc[1, "记录数"])
chk("Y1+Y2 = 分析口径", n1 + n2, N_ANA)
NET = n2 - n1
chk("净增长", NET, 362541)

t06 = L("t06_yoy_by_type")
w3 = int(t06[t06["问题类型"].isin(["Snow or Ice", "HEAT/HOT WATER", "Street Condition"])]["增减量"].sum())
chk("天气三类增量", w3, 167706)
chk("天气三类占净增长%", f"{100*w3/NET:.1f}", "46.3")
chk("t06 占净增长% 自洽", f"{100*t06.loc[0,'增减量']/NET:.1f}", f"{t06.loc[0,'占净增长%']:.1f}")

t07 = L("t07_yoy_by_agency")
chk("t07 增减量合计 = 净增长", int(t07["增减量"].sum()), NET)

t16 = L("t16_nypd_outcome")
chk("NYPD 分类占比合计≈100（容许四舍五入）",
    abs(t16["占比%"].sum() - 100) < 0.05, True)
chk("NYPD 采取行动%", t16[t16["结案性质"].str.startswith("A")]["占比%"].iloc[0], 27.96)
chk("NYPD 未发现问题%", t16[t16["结案性质"].str.startswith("B")]["占比%"].iloc[0], 65.84)

t19 = L("t19_dup_overall")
chk("簇数+冗余 = 带地址记录数",
    int(t19["去重后事件簇数"].iloc[0] + t19["冗余记录数"].iloc[0]),
    int(t19["可判定记录数(有地址)"].iloc[0]))
chk("冗余占分析口径全量%",
    f"{100*t19['冗余记录数'].iloc[0]/N_ANA:.2f}", f"{t19['冗余占分析口径全量%'].iloc[0]:.2f}")

t23 = L("t23_borough")
chk("行政区记录数 + Unspecified = 分析口径", int(t23["记录数"].sum()) + 6331, N_ANA)
chk("行政区占比合计%", f"{t23['占全市%'].sum():.0f}", "100")

t30 = L("t30_noise_yoy_effect")
chk("t30 含聚集 - 聚集 = 分析口径 (Y1)",
    int(t30.loc[0, "住宅噪音（含异常聚集）"] - t30.loc[0, "其中异常聚集"]),
    int(t30.loc[0, "住宅噪音（分析口径）"]))
chk("住宅噪音同比 含聚集%",
    f"{100*(t30.loc[1,'住宅噪音（含异常聚集）']-t30.loc[0,'住宅噪音（含异常聚集）'])/t30.loc[0,'住宅噪音（含异常聚集）']:.1f}",
    "-2.8")
chk("住宅噪音同比 分析口径%",
    f"{100*(t30.loc[1,'住宅噪音（分析口径）']-t30.loc[0,'住宅噪音（分析口径）'])/t30.loc[0,'住宅噪音（分析口径）']:.1f}",
    "7.5")

t27 = L("t27_top_types")
chk("t27 记录数降序且累计递增",
    bool(t27["记录数"].is_monotonic_decreasing and t27["累计占比%"].is_monotonic_increasing), "True")
chk("t27 前10累计", f"{t27.loc[9,'累计占比%']:.1f}", "55.8")

print("\n--- 关键数字回查 ---")
bad = 0
for name, got, exp in checks:
    ok = got == exp
    bad += 0 if ok else 1
    print(f"{'OK ' if ok else '!! '}{name:<34s} 实际={got:<14s} 期望={exp}")
print("\n全部一致" if bad == 0 else f"\n{bad} 处不一致，需检查")
