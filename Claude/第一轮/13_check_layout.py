"""步骤8（续）：检查合并单元格是否重叠、是否与已有数据冲突（Excel 打开报错的常见原因）。"""
import openpyxl, pathlib
from openpyxl.utils import range_boundaries

wb = openpyxl.load_workbook(pathlib.Path(r"D:\项目2\Claude\最终成果.xlsx"))
bad = 0
for s in wb.sheetnames:
    ws = wb[s]
    cells = {}
    for rng in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(rng))
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                if (rr, cc) in cells:
                    print(f"!! {s} 合并重叠 {rng} 与 {cells[(rr,cc)]}")
                    bad += 1
                cells[(rr, cc)] = str(rng)
        # 合并区内除左上角外不应有值
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                if (rr, cc) == (r1, c1):
                    continue
                if ws.cell(rr, cc).value not in (None, ""):
                    print(f"!! {s} 合并区 {rng} 内 ({rr},{cc}) 有值 {ws.cell(rr,cc).value!r}")
                    bad += 1
    print(f"{s:<18s} 合并区 {len(ws.merged_cells.ranges)} 个")
print("\n布局检查通过，无重叠/冲突" if bad == 0 else f"\n发现 {bad} 处问题")

# ---- 图表数据区校验：每个系列的取值范围必须落在表格数据行内且非空 ----
print("\n--- 图表数据区校验 ---")
import zipfile, re
bad2 = 0
with zipfile.ZipFile(r"D:\项目2\Claude\最终成果.xlsx") as z:
    charts = sorted(n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
    for cn in charts:
        xml = z.read(cn).decode("utf-8")
        title = (re.search(r"<c:tx>.*?<a:t>(.*?)</a:t>", xml, re.S) or [None, "(无标题)"])[1]
        refs = re.findall(r"<c:val>.*?<c:f>(.*?)</c:f>", xml, re.S)
        cats = re.findall(r"<c:cat>.*?<c:f>(.*?)</c:f>", xml, re.S)
        for vf, cf in zip(refs, cats):
            def span(ref):
                m = re.search(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", ref)
                return int(m.group(4)) - int(m.group(2)) + 1 if m else -1
            sheet = vf.split("!")[0].strip("'")
            nv, nc = span(vf), span(cf)
            # 取值区首行不应是表头（表头为文本，数据为数字）
            ws2 = wb[sheet]
            m = re.search(r"\$?([A-Z]+)\$?(\d+):", vf)
            col = openpyxl.utils.column_index_from_string(m.group(1))
            first = ws2.cell(int(m.group(2)), col).value
            hdr = ws2.cell(int(m.group(2)) - 1, col).value
            okrow = isinstance(first, (int, float)) and isinstance(hdr, str)
            if nv != nc or not okrow:
                print(f"!! {cn} [{title}] 取值{nv}点/类别{nc}点 首格={first!r} 上一格={hdr!r}")
                bad2 += 1
print(f"图表 {len(charts)} 个，"
      + ("数据区校验通过" if bad2 == 0 else f"发现 {bad2} 处问题"))
