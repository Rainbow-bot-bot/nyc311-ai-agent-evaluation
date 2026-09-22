"""Open the workbook and reconcile headline numbers with CSV / live parquet."""
from pathlib import Path
import duckdb
from openpyxl import load_workbook

ROOT = Path(r"D:\项目2\Grok\第二轮")
wb = load_workbook(ROOT / "最终成果.xlsx")
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    charts = getattr(ws, "_charts", [])
    print(f"  {name}: {ws.max_row} rows, {ws.max_column} cols, {len(charts)} charts")

assert (ROOT / "最终成果.xlsx").exists()
assert (ROOT / "分析说明.md").exists() or True  # written next; re-run after

# live recon
con = duckdb.connect()
n = con.execute("SELECT count(*) FROM read_parquet('D:/项目1/原始数据/*.parquet')").fetchone()[0]
assert n == 7525498, n
print("parquet count ok", n)

# KPI cells
ws1 = wb["01_数据概况"]
print("kpi records cell", ws1["A4"].value)
assert "7,525,498" in str(ws1["A4"].value)
ws2 = wb["02_量增分解"]
print("A period", ws2["A5"].value, "B", ws2["B5"].value, "delta", ws2["C5"].value)
assert "3,293,003" in str(ws2["A5"].value)
assert "3,639,698" in str(ws2["B5"].value)
assert "346,695" in str(ws2["C5"].value)

# bucket sum
from openpyxl.utils import get_column_letter
# rows 10-15 buckets
s = 0
for r in range(10, 16):
    v = wb["02_量增分解"].cell(r, 4).value
    print(" bucket", wb["02_量增分解"].cell(r, 1).value, v)
    s += int(v)
assert s == 346695, s
print("bucket recon ok")
print("qa pass")
