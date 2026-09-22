from pathlib import Path
import json, math
import openpyxl,win32com.client,pymupdf as fitz
R=Path(__file__).parent;E=R/'evidence';P=E/'previews';P.mkdir(exist_ok=True)
path=R/'最终成果.xlsx'
# Native engine: no repair option, update no external links, full recalculation.
app=win32com.client.DispatchEx('Excel.Application');app.Visible=False;app.DisplayAlerts=False
book=None
try:
    book=app.Workbooks.Open(str(path),UpdateLinks=0,ReadOnly=False)
    app.CalculateFullRebuild()
    charts=[]
    for sheet in book.Worksheets:
        for chart in sheet.ChartObjects():
            charts.append({'sheet':sheet.Name,'name':chart.Name,'series':[x.Formula for x in chart.Chart.SeriesCollection()]})
    book.Save()
    views={'总览':'A1:L40','月度趋势':'A1:L32','同比拆解':'A1:L38','峰值核查':'A1:H41','渠道变化':'A1:G26','数据质量':'A1:D23','口径与复现':'A1:B23'}
    for idx,(name,area) in enumerate(views.items(),1):
        sh=book.Worksheets(name);old=sh.PageSetup.PrintArea;old_tall=sh.PageSetup.FitToPagesTall
        sh.PageSetup.PrintArea=area;sh.PageSetup.FitToPagesWide=1;sh.PageSetup.FitToPagesTall=1
        sh.ExportAsFixedFormat(0,str(P/f'{idx:02d}.pdf'))
        sh.PageSetup.PrintArea=old;sh.PageSetup.FitToPagesTall=old_tall
    book.Close(SaveChanges=False);book=None
finally:
    if book is not None:book.Close(SaveChanges=False)
    app.Quit()
for pdf in P.glob('*.pdf'):
    doc=fitz.open(pdf);page=doc[0];page.get_pixmap(matrix=fitz.Matrix(1.7,1.7)).save(pdf.with_suffix('.png'));doc.close()
v=openpyxl.load_workbook(path,data_only=True);f=openpyxl.load_workbook(path,data_only=False)
errors=[];formula_count=0
for sh in f:
    for row in sh:
        for cell in row:
            if cell.data_type=='f':
                formula_count+=1;cached=v[sh.title][cell.coordinate]
                if cached.value is None or cached.data_type=='e':errors.append([sh.title,cell.coordinate,cached.value])
assert not errors,errors[:20]
assert v['月度趋势']['B14'].value==2378736
assert v['月度趋势']['C14'].value==2666078
assert v['月度趋势']['D14'].value==287342
assert v['月度趋势']['B46'].value==7525498
assert v['总览']['B7'].value==2378736 and v['总览']['C7'].value==2666078
st=json.loads((E/'standardization_summary.json').read_text())
assert math.isclose(v['渠道变化']['B20'].value,st['base_common_rate'],abs_tol=1e-12)
assert math.isclose(v['渠道变化']['B21'].value,st['fixed_mix_current_rate'],abs_tol=1e-12)
assert len(charts)==2 and all(len(x['series']) for x in charts)
result={'opened_in':'Microsoft Excel 16.0','full_recalculation':True,'sheets':v.sheetnames,'formula_count':formula_count,'formula_errors':errors,'charts':charts,'native_pdf_previews':list(views),'headlines_checked':True,'xlsx_bytes':path.stat().st_size}
(E/'workbook_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
