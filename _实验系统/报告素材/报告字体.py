"""明确指定中文和西文字体，清除模板主题字体的继承。"""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_font(rpr,heading=False):
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    rf.attrib.clear()
    for key,value in {'ascii':'Arial' if heading else 'Times New Roman','hAnsi':'Arial' if heading else 'Times New Roman','eastAsia':'黑体' if heading else '宋体','cs':'Arial' if heading else 'Times New Roman'}.items():rf.set(qn('w:'+key),value)
    lang=rpr.find(qn('w:lang'))
    if lang is None:lang=OxmlElement('w:lang');rpr.append(lang)
    lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'zh-CN')

def normalize_fonts(doc):
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        st=doc.styles[name];set_font(st.element.get_or_add_rPr(),name=='Title' or name.startswith('Heading'))
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    paragraphs=list(doc.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs);paragraphs.extend(section.footer.paragraphs)
    for p in paragraphs:
        for border in list(p._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
        heading=p.style.name=='Title' or p.style.name.startswith('Heading')
        for r in p.runs:set_font(r._element.get_or_add_rPr(),heading)
    for tag in ['hideSpellingErrors','hideGrammaticalErrors']:
        if doc.settings.element.find(qn('w:'+tag)) is None:doc.settings.element.append(OxmlElement('w:'+tag))
