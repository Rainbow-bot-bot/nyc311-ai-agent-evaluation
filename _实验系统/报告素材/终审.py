"""只读核对现行S评分、缓存总分、文件结构及阅读入口。"""
from pathlib import Path
import hashlib,json,re,sys
from collections import defaultdict
from urllib.parse import unquote

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
import 评测数据处理 as pipeline

def main():
    from openpyxl import load_workbook
    scores=pipeline.read_current_scores()
    result=pipeline.verify_current_scores(scores=scores)
    totals=defaultdict(float);dims=defaultdict(float)
    for entry in scores.values():
        s=entry['评分原文'];totals[s['AI'],s['轮次']]+=s['得分'];dims[s['AI'],s['轮次'],s['条款'][0]]+=s['得分']
    file=ROOT/'交付成果/项目2_AI评测分析.xlsx'
    before=hashlib.sha256(file.read_bytes()).hexdigest()
    book=load_workbook(file,data_only=True)
    try:
        assert len(book.sheetnames)==10 and sum(len(s._charts) for s in book)==8
        for row in list(book['运行记录'].values)[1:]:
            if not row[0]:continue
            key=(row[0],1 if row[1]=='第一轮' else 2)
            expected=totals.get(key)
            assert row[13]==expected,(key,row[13],expected)
        for name,cells in {'评分细项':{'B2'},'分维度比较':{'B3','F3'},'单AI分析':{'C3','G3'},'评测总览':{'C3'}}.items():
            assert cells.issubset({str(v.sqref) for v in book[name].data_validations.dataValidation}),name
        assert book['字段说明']['I10'].value==book['字段说明']['J10'].value=='kimi-k3'
        assert book['字段说明']['I8'].value==book['字段说明']['J8'].value=='gemini-3.8-flash'
    finally:book.close()
    systems=['GPT','Grok','DeepSeek','GLM','Qwen','Gemini']
    result['六系统配对']={ai:[totals[ai,1],totals[ai,2]] for ai in systems}
    result['六系统均值']=[sum(totals[a,r] for a in systems)/6 for r in (1,2)]
    result['平均增分']=result['六系统均值'][1]-result['六系统均值'][0]
    delta=sum(totals[a,2]-totals[a,1] for a in systems)
    result['调查维度贡献']=sum(dims[a,2,'C']-dims[a,1,'C'] for a in systems)/delta
    for rel in ['README.md','资料/复现说明.md','资料/评判标准.md','资料/公开仓库说明.md','交付成果/开始阅读.html','交付成果/查证.html']:
        p=ROOT/rel;text=p.read_text(encoding='utf-8-sig')
        links=re.findall(r'(?:href|src)="([^"]+)"',text)+re.findall(r'\]\(([^)]+)\)',text)
        for link in links:
            if link.startswith(('#','http','javascript:','data:')):continue
            assert (p.parent/unquote(link.split('#')[0])).exists(),(rel,link)
    nb=json.loads((ROOT/'交付成果/01_评测数据整理与验收.ipynb').read_text(encoding='utf-8'))
    for cell in nb['cells']:
        if cell['cell_type']=='code':
            assert cell['execution_count'] is not None
            assert all(o['output_type']!='error' for o in cell.get('outputs',[]))
    assert hashlib.sha256(file.read_bytes()).hexdigest()==before
    result['范围']='S分项与依据同步、缓存总分、下拉定义、阅读链接和Notebook已保存输出；原生交互另运行终审交互.py'
    print(json.dumps(result,ensure_ascii=False,indent=2))
    return result

if __name__=='__main__':main()
