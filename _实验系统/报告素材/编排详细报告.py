from pathlib import Path
import json,hashlib,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(r'D:\项目2'); D=R/'交付成果'; A=R/'_实验系统'/'报告素材'
cases=json.loads((A/'各AI分析.json').read_text(encoding='utf-8'))
doc=Document()
s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.9);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
for name,size,chinese,bold in [('Normal',12,'宋体',False),('Title',24,'黑体',False),('Heading 1',16,'黑体',False),('Heading 2',13,'黑体',False),('Heading 3',11,'黑体',False),('Caption',9,'宋体',False)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),chinese)
    st.paragraph_format.line_spacing=1.25;st.paragraph_format.space_after=Pt(8)
    for tag in ['numPr','pBdr']:
        for el in list(st._element.iter(qn('w:'+tag))):el.getparent().remove(el)
doc.styles['Normal'].paragraph_format.widow_control=True
header=s.header.paragraphs[0];header.text='AI 数据分析 Agent 自主交付评测';header.alignment=2
header.style=doc.styles['Caption']
footer=s.footer.paragraphs[0];footer.alignment=1
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
settings=doc.settings.element
proof=OxmlElement('w:proofState');proof.set(qn('w:spelling'),'clean');proof.set(qn('w:grammar'),'clean');settings.append(proof)
md=['# AI 数据分析 Agent 自主交付评测报告\n']; manifest=[]

def p(t):
    x=doc.add_paragraph(t);md.append(t+'\n')
    return x
pending_page=False
def h(t,l=1):
    global pending_page
    x=doc.add_heading(t,l)
    if pending_page:x.paragraph_format.page_break_before=True;pending_page=False
    md.append('#'*(l+1)+' '+t+'\n');return x
def page():
    global pending_page
    pending_page=True
def fig(name,caption,source,width=17):
    kind='表' if caption.startswith('表') else '图'
    number=1+sum(x['caption'].startswith(kind) for x in manifest)
    caption=re.sub(r'^[表图]\d+',kind+str(number),caption)
    file=A/(name+'.png');pp=doc.add_paragraph();pp.alignment=1
    pp.add_run().add_picture(str(file),width=Cm(width));pp.paragraph_format.keep_with_next=True
    pp.paragraph_format.space_after=Pt(3)
    cp=doc.add_paragraph(caption,'Caption');cp.alignment=1
    md.append(f'![{caption}](../_实验系统/报告素材/{name}.png)\n\n来源：{source}。\n')
    manifest.append({'file':file.name,'caption':caption,'source':source,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})
def lead(label,text):
    x=doc.add_paragraph();x.add_run(label+'　').bold=True;x.add_run(text);md.append('**'+label+'** '+text+'\n')

doc.add_paragraph('AI 数据分析 Agent\n自主交付评测报告','Title')
p('NYC311 数据分析任务　2026年9月21日')
h('摘要')
p('本项目考察个人使用AI完成数据分析的实际效果。评测限于两轮，比较结果是否可靠、费用与耗时是否值得，以及使用者还需承担多少续跑、接管和修正工作。失败记录全部保留，不追加试验后挑选最好的一次。')
p('本评测要求八个AI系统基于同一份纽约311服务请求数据，自主选题、开展分析并提交Excel及分析说明。GPT、Grok、DeepSeek、GLM、Qwen和Gemini完成了两轮独立运行。Claude第一轮完成主要调查，最终由Grok接手完成交付；Kimi第一轮多次提前结束，第二轮复测也未交付成品。')
p('各系统的主要差异集中在异常调查和结论核验。GPT继续检查雪冰峰日与渠道构成，Grok追查供暖集中的地块，DeepSeek检查高集中地址和固定观察期。GLM、Qwen增加了贡献分解和异常检查；Gemini扩大了调查范围，但仍误读了部分结案文本。')
p('六个系统第二轮运行期间都没有收到追加指令，交付后却仍有10项登记修正。已完成的交付仍存在内容错误，包括时间转述、字段数量、比较月份，以及关键词在原句中的含义。')
h('研究问题',2)
for t in ['1. 各系统能否在无人持续纠偏的情况下完成交付？','2. 成果质量的差距来自哪些维度和具体工作？','3. 各系统能否发现并修正自己的错误？','4. 第二轮方法提示条件下，哪些方面改善，哪些问题仍然存在？','5. 将质量、时间、成本和后续修正放在一起，如何选择？']:p(t)

page();h('1 评测设计')
h('1.1 任务与样本',2)
p('评测安排两轮，比较各系统的分析质量、运行费用、用时及所需人工干预。Claude第二轮未形成独立运行记录。费用按API目录价估算，人工干预以已登记任务描述，未记录实际账单和人工工时。')
p('任务材料包括25个Parquet分片，共7,525,498条NYC311服务请求、44个字段。各系统自行选择分析问题，最终提交Excel和分析说明。第一轮使用基础提示，第二轮加入Evidence-led方法提示，Grok和Gemini还收到额外方法清单。模型、运行工具和提示条件共同影响结果，每种条件只运行一次。')
p('两轮配对比较采用GPT、Grok、DeepSeek、GLM和Qwen五个系统。Gemini第二轮越界搜索项目根目录，且模型子版本缺少可核记录，未通过G2/G4；其成品及评分单列为观察记录。Claude的混合交付及Kimi的退出、失败记录单列分析，未评分的项目留空。第二轮没有可验证的Claude独立运行。')
h('1.2 评分与核验',2)
p('评分包括成果质量Q和过程可靠性R。Q评价分析内容与最终成品，R评价独立执行、纠错过程及完成声明的真实性。两者各满分100，分别报告。')
p('Q包含五个维度、21个细项：数据判断20分，检查数据理解和选题；证据计算25分，检查数字、比较条件与解释；调查推进30分，检查拆解、追查和对照；最终交付15分，检查文件、说明及终稿核验；复杂度与范围10分，检查方法和投入是否合适。')
p('R包含五项：独立承担25分，自我纠错25分，路线控制20分，完成真实性20分，协议纪律10分。正文说明主要得失分依据，附表列出每个细项的两轮得分和满分。评分版本及原始记录入口见附录。')

page();h('1.3 模型与运行条件')
fig('运行条件','表1  模型、思考强度与Harness','字段说明 H1:L14',17)
p('本项目借助AI完成数据整理、分析与报告编制。')
p('Harness指承载模型、组织工具调用及会话执行的环境。')
p('Claude第一轮因Cursor宿主额度限制切换到Grok收尾，Q只评价混合成品，独立过程R留空。第二轮没有可验证的独立运行，因此各项及轮间变化留空。')
p('Kimi第一轮连续提前结束，依赖人工续跑后退出；第二轮真实复测反复读取同一清单，评测方在确认循环后终止。两轮均未交付完整成品，未评分。')

qnames=['数据判断','证据计算','调查推进','最终交付','复杂度与范围']
for i,(ai,c) in enumerate(cases.items(),2):
    page();h(f'{i} {ai}')
    h('第一轮的工作',2);p(c['first'])
    h('第二轮的工作',2);p(c['second'])
    start=51+list(cases).index(ai)*14
    fig(ai+'维度',f'表{7+list(cases).index(ai)}  {ai}两轮维度得分',f'分维度比较 A{start}:E{start+(12 if ai in ("Qwen","Gemini") else 11)}',16.8)
    page();h(f'{i}.1 {ai}评分依据')
    for label,txt in zip(qnames,c['q']):
        lead(label,txt)
    h('过程可靠性',2);p(c['r'])
    p(f'{ai}的完整细项分数见附表{i-1}。')

page();h('8 Claude')
h('第一轮的调查与交付',2)
p('Claude第一轮核对字段与数据范围，调查高频地址、增长来源、类别变化、关闭时长、结案原句和重复记录。它保存30张中间汇总表，生成13个工作表。按“地址×类别单日峰值不少于200条”筛出4个住宅噪声组合，共185,791条，再比较剔除前后的趋势。它还用日均值处理365天与364天的窗口差异，检查成熟队列、机构时长和NYPD结案描述。')
p('第一轮因Cursor宿主额度限制，运行从Claude Opus 5切换到Grok。记录中Claude有51次工具请求，切换后的Grok有31次。最终重建、说明对齐和部分文件读回由Grok完成，因此这份成品按混合交付评价，不给Claude单独的过程可靠性分数。模型分段Token也缺少可核实的记录。')
fig('Claude维度','表13  Claude第一轮混合成品维度得分','分维度比较 A135:E147',16.8)

page();h('8.1 Claude评分依据')
lead('数据判断','Claude基本理解请求粒度、44个字段和全量范围，也选出了6个具体问题。不过，它把高频聚集当成必须删除的无效记录，并在标题中混淆最大单地址与4个组合总量，数据判断得17/20。高频记录可以用来做剔除前后的对照，删除前仍需证明它们确实无效。')
lead('证据计算','证据计算得16/25。总量、窗口计数和既定去重规则下的冗余量可以复算，但前20类的增量与全量增量未对齐，自检又部分使用常量和旧汇总。365天与364天的差异虽已说明，末日不完整和时效观察时间仍影响比较。“真实重复只会更多”、天气已排除其他解释、官方改名等说法缺少依据，解释边界仅得1/5。')
lead('调查推进','增量拆解、深入调查均得满分，调查推进合计25/30。三类投诉贡献的167,706条可以与362,541条总增量对照，地址日峰、相邻聚集、前后趋势和结案原句也查得较细。剩余扣分涉及异常剔除是否有效、去重是否误合并，以及成熟样本如何选择。')
lead('最终交付','最终交付得14/15。数字、图表和计算过程可以回查，但增量加总声明的错误未被检查发现，最终文件检查得3/4。Grok也参与了最终核验。后来迁移目录造成的旧路径问题由评测方维护，不计入Claude失分。')
lead('复杂度与范围','复杂度与范围得9/10。DuckDB聚合和统一保存函数都用于实际分析；扣分来自几次写入被截断后继续重试、再拆分的返工，资源判断为2/3。缺少Token记录，无法估算这部分费用。')
h('过程与比较',2)
p('Claude与第一轮Qwen的成品质量同为81分。Claude调查推进多4分，证据计算少3分，复杂度与范围少1分。Claude完成了主要调查，Grok参与了交付，两者共同形成最终结果；独立承担、自我纠错、路线控制、完成真实性和协议纪律五项R评分均留空。完整细项见附表7。')

page();h('9 Kimi')
h('第一轮为何退出',2)
p('Kimi第一轮连续提前结束，需要人工续跑，后来退出评测。参与登记和后续规划记录了这一经过；现有目录只有任务提示，没有最终Excel和分析说明。第一轮耗时、Token及续跑次数缺少原始记录，保留空白。')
h('第二轮补充复测',2)
p('第二轮先遇到启动器未关闭标准输入的问题，当时模型尚未真正启动。这次基础设施失败不计为模型失败。修复后，Kimi反复读取同一个data_manifest.json，即使运行工具已返回file_unchanged或Wasted call仍继续读取。评测方确认循环后停止运行，未得到Excel和分析说明。')
p('真实复测用时约158.7秒，即2.65分钟，按API目录价估算约1.48元。14个实质动作中13个没有收益，占92.86%。另一份早期计时记录为171.252秒，本报告采用原生会话时间。')
h('评分处理',2)
lead('数据判断与证据计算','两轮都没有完整成品，第二轮主要停在反复读清单。数据判断和证据计算缺少评分材料，留空。')
lead('调查推进与最终交付','第一轮提前结束，第二轮陷入重复读取，两次都未完成后续调查和交付。调查推进、最终交付同样不评分。')
lead('复杂度与过程可靠性','第一轮需要人工续跑，第二轮未根据“文件未变化”的反馈调整做法，运行连续性和路线控制有明确问题。其他过程项目缺少足够记录，R不发布。')

page();h('10 成果质量比较')
fig('Q1','表1  第一轮成果质量分项','分维度比较 A1:H10',16.8)
fig('Q2','表2  第二轮成果质量分项','分维度比较 A13:H22',16.8)
page();h('10.1 分项差距与同分差异')
lead('数据判断','第二轮数据判断差距很小：GPT、Grok和DeepSeek均为20分，GLM、Qwen和Gemini均为19分。这些交付都已检查基本数据范围。GLM把44个字段写成43个，是其中一个具体失分点。')
lead('证据计算','正式比较的五个系统中，证据计算从Qwen的21分到GPT的25分，相差4分；Grok24分、GLM23分、DeepSeek22分。纳入Gemini的16分观察评分后，六系统相差9分。DeepSeek写错终稿时间，Qwen的结构解释与比较月份有问题，Gemini误读结案文本；这些具体错误解释了部分分差。')
lead('调查推进','GPT、Grok和DeepSeek的调查推进均为30分，GLM和Qwen为28分，Gemini为23分。Grok追到地块集中度，DeepSeek追到地址与关闭方式，前两者都加入了针对性的对照。Gemini也展开了更多调查，但仍保留缺少依据的原因解释，适时停止为0/3。')
lead('最终交付','最终交付从GPT的15分到Gemini的12分。Grok、GLM和Qwen各14分，DeepSeek13分。它们都做了文件检查，差别在于是否查到终稿中的文字错误，以及检查脚本本身是否有效。')
lead('复杂度与范围','GPT、Grok、DeepSeek和GLM均为10分，Qwen、Gemini为8分。Qwen主要受较重的重建和核对投入影响，Gemini还有范围与未经授权信息的问题。')
p('GPT质量100分、Grok98分，差异来自证据计算和最终交付各1分。Qwen分项合计90，主线时间窗错误使最终Q封顶59；GLM为94分。DeepSeek比GLM总分只多1分，其中数据判断多1分、调查推进多2分，证据计算和最终交付各少1分。')

page();h('11 过程可靠性比较')
fig('R1','表3  第一轮过程可靠性分项','分维度比较 A25:H34')
fig('R2','表4  第二轮过程可靠性分项','分维度比较 A37:H46')
lead('独立承担与协议纪律','正式五系统两项均为满分。Gemini观察评分分别为24/25和8/10，协议纪律扣分来自越界搜索目录；未发现读取他人成果或评分。Claude与Kimi不纳入独立过程比较。')
lead('自我纠错','Grok得25分，GPT24分，Gemini16分。Qwen较第一轮提高4分，主要修复文件和图表；Gemini提高3分，关键语义与原因解释仍需外部指出。')
lead('路线控制','Grok、Qwen和Gemini得20分，其余三个系统19分。Gemini按计划完成文件，文本解释中的错误仍留在终稿。')
lead('完成真实性','GPT20分，Gemini13分，差距最大。Grok因恒真的文件存在检查扣1分，与自我纠错增分抵消，R仍为99。')

page();h('12 两轮变化')
fig('分项增量','表5  五个配对系统的分项变化与封顶调整','评测总览 B29:N38')
p('五个系统最终Q的变化分别为GPT +7、Grok +4、DeepSeek +6、GLM +8、Qwen −22。均值由88.6升至89.2，增加0.6分，中位变化为6分。Gemini第二轮未通过Gate，不进入配对汇总。')
p('封顶前，五系统分项合计增加34分：数据判断4分、证据计算5分、调查推进26分、最终交付−1分，复杂度与范围不变。调查推进占分项增量76.5%。')
p('Qwen分项合计90，但11个月与9个月的关闭时长比较支撑了主要发现3，按冻结规则封顶59，调整−31分。分项提升与最终质量下降同时存在：更多工作没有解决主线比较错误。')
p('GPT的深入调查和修正判断经本次重审分别改为8/8、4/4。原98分缺少两处分项扣分依据，现Q为100。旧分数与本次裁定同时保留。')
h('12.1 评分变化与实际动作',2)
p('增加的调查可以在操作记录中找到：GPT排除峰日并固定渠道构成，Grok剔除固定前10个地块，DeepSeek比较高集中地址和固定观察期，GLM补上增量分解，Qwen增加异常调整和复现检查。这些新增动作与调查推进得分的提高相对应。')

page();h('13 耗时、费用与后续修正')
fig('总览-第二轮','表6  第二轮交付与资源指标','评测总览 B8:N16；第二轮')
fig('总览图2','图1  第二轮原生耗时','评测总览 图表2；第二轮',10.5)
fig('总览图3','图2  第二轮API目录价低档','评测总览 图表3；第二轮',11)
p('第二轮GLM约9.6分钟，目录价6.48元；DeepSeek约21.7分钟，1.37—2.74元；GPT约17.8分钟，38.96元；Qwen约60.2分钟，14.33元。Grok和Gemini缺少Token记录，费用留空。图中显示DeepSeek区间下限，表中保留完整区间。金额按API目录价估算，与实际账单可能不同。')
p('本次GLM比Qwen用时更短、费用更低，质量与过程可靠性得分也更高。GPT交付后未登记修正任务，但可估算费用最高；DeepSeek费用最低，仍有1项终稿修正。')

page();h('13.1 无收益动作与后续修正')
fig('总览图4','图3  第二轮无收益动作占比','评测总览 图表4；第二轮',16)
p('第二轮六个独立系统共留下362条可查动作，其中300条属于实质动作，5条被判定没有收益。GPT为2/32，DeepSeek为1/88，GLM为2/41，Grok、Qwen和Gemini分别为0/38、0/62、0/39。无收益动作占比AWR排除管理动作；Grok和Gemini缺少逐动作时长，不能再比较这些动作占用了多少时间。')
p('运行期间六者都没有收到追加指令。交付后登记修正共10项：GPT 0项，Grok、DeepSeek、GLM各1项，Qwen 2项，Gemini 5项。构成与分类解释占3项，时间与比较范围占2项，解释缺证据占2项，数值或对象错配、记录与现实对象混淆、文本语义各1项。')
p('各项修正的工作量差别很大。GLM的字段数量可以局部改正，Gemini的文本语义和天气解释则可能影响整段结论。项目未记录修正工时。')
p('Qwen和Gemini没有动作被标为无收益，交付后仍各有2项、5项登记修正。Qwen检查了文件和图表，Gemini检查了文件完整性，两者都遗漏了最终解释中的问题。')

conclusions=json.loads((A/'评测结论.json').read_text(encoding='utf-8'))
page();h('14 研究结论')
p(conclusions['总判断'])
h('14.1 文件完成后仍需人工修正',2)
p('五个独立系统的成品仍需修正，主要问题涉及数值转述、统计范围、对象定义与文本语义。它们通常能发现脚本或文件故障，对解释是否成立、终稿是否忠实于计算结果的检查较弱。GPT本次未登记交付后修正；其他系统的检查次数或无收益动作占比，均不足以直接判断成品准确性。')
h('14.2 分差主要来自计算和解释',2)
p('第二轮正式比较的五个系统，数据判断相差1分，证据计算相差4分；含Gemini观察评分时，证据计算相差9分。各系统普遍能识别数据范围，差异更多体现在如何选择分母、对齐窗口、理解字段和限制原因解释。Grok与DeepSeek调查较深入，GPT在计算与终稿核验上更完整，Gemini的关键弱点是语义和因果判断。总分相近的系统，复核重点并不相同。')
h('14.3 第二轮增加了调查，终稿检查改善有限',2)
p(conclusions['方法判断'])
page();h('14.4 各系统综合评价')
for ai,judgement in conclusions['系统评价'].items():lead(ai,judgement)


page();h('附录 评分裁定与资料说明')
p('冻结规则保持不变。原评分版本为第一轮1.0、第二轮enhanced_review_1.0，Claude为claude_assisted_quality_20260918_v1。本次final_review_20260921记录GPT分项重审、Qwen封顶与Gemini样本资格，原始评分保留。Excel“评分与用量”保留细项分数和证据编号，“核验记录”保留问题、对照值及归属，“行为记录”保留动作次序。')
p('报告全部图表取自交付目录中的主Excel。主表支持按AI和轮次查看结果，“分维度比较”用于横向比较，“评分细项”列出以下七份附表。Claude缺少独立过程评分，Kimi缺少完整评分，空白均保留。')
p('在“查证.html”输入证据编号可查看已有核验；Notebook保留数据整理、两轮配对和分维度计算。')
p('本评测由一名评审评分，未测量评审间一致性，也未记录人工监督和修正工时。')


detail_sources=json.loads((A/'评分细项来源.json').read_text(encoding='utf-8'))
for i,b in enumerate(detail_sources,1):
    page();h(f'附表{i} {b["ai"]}评分细项')
    fig(b['ai']+'细项',f'表0  {b["ai"]}完整细项得分','评分细项 '+b['range'],17)

from 报告字体 import normalize_fonts
normalize_fonts(doc)
out=D/'项目2_AI数据分析Agent自主交付评测报告.docx';doc.save(out)
(D/'项目2_自主交付研究报告.md').write_text('\n'.join(md),encoding='utf-8')
(A/'图表来源.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('详细报告已写入',out,'图片',len(doc.inline_shapes),'正文段落',len(doc.paragraphs))
