from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from 报告字体 import normalize_fonts
import json,hashlib
A=Path(__file__).parent
d=Document();s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.7);s.bottom_margin=Cm(1.7);s.left_margin=Cm(2);s.right_margin=Cm(2)
for name,size in [('Normal',11.5),('Title',21),('Heading 1',16),('Heading 2',12.5),('Caption',9)]:
 st=d.styles[name];st.font.size=Pt(size);st.font.bold=False;st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.line_spacing=1.22;st.paragraph_format.space_after=Pt(8)
header=s.header.paragraphs[0];header.text='AI 数据分析 Agent 自主交付评测  简洁版';header.style='Caption';header.alignment=2
foot=s.footer.paragraphs[0];foot.alignment=1
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
def p(t):return d.add_paragraph(t)
def h(t,new=False):
 x=d.add_heading(t,1);x.paragraph_format.page_break_before=new;return x
def lead(label,t):
 x=d.add_paragraph();x.add_run(label+'　').bold=True;x.add_run(t)
manifest=[]
def fig(name,caption,source,width=15):
 file=A/(name+'.png');x=p('');x.alignment=1;x.add_run().add_picture(str(file),width=Cm(width));x.paragraph_format.space_after=Pt(2);x.paragraph_format.keep_with_next=True
 y=d.add_paragraph(caption,'Caption');y.alignment=1
 manifest.append({'file':file.name,'caption':caption,'source':source,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})

d.add_paragraph('AI 数据分析 Agent\n自主交付评测简洁版','Title')
p('NYC311 数据分析任务　2026年9月21日')
p('八个系统使用相同任务材料，自主选题并提交Excel与分析说明。数据包含25个Parquet分片、7,525,498条服务请求和44个字段。第一轮使用基础提示，第二轮增加Evidence-led方法提示，Grok、Gemini另有方法清单。')
p('评测面向个人使用，限于两轮。除分析质量外，还比较费用、耗时，以及续跑、接管和修正工作。')
h('1 各系统完成的主要工作')
lead('GPT','第一轮分解类别、行政区增长，第二轮进一步核查雪冰峰日和渠道构成。最高5天占同期雪冰请求的59.42%；排除峰日并固定渠道构成后继续比较，最终检查Excel公式和图表。')
lead('Grok','第一轮按互斥类别分解全市增量，第二轮检查供暖地块集中度、积雪高量日及坑洞。排除固定前10个地块后，供暖请求仍增长21.96%。')
lead('DeepSeek','两轮均调查异常地址和关闭时长，第二轮增加邮编集中度、固定30天关闭率及数据重建。第一轮固定机构与类型构成后，冬季时长倍数由3.8667降至1.1537。')
lead('GLM','第一轮重点追查道路异常，第二轮增加同比增量分解、住宅噪声与热点地址分析，并修正月均值和关闭率的统计月份。')
lead('Qwen','第一轮完成类别、地区和时长比较，第二轮增加噪声尖峰、渠道及异常调整分析，多次重建工作簿并核对图表与说明。')
lead('Gemini','第一轮分析异常地址、供暖、违停及机构时长，第二轮完善日期与字段检查，扩大部门、类别和处置结果分析，但结案文本解释仍有错误。')
lead('Claude','第一轮保存30张中间汇总表、生成13个工作表，调查地址集中、重复规则与结案描述。Claude因Cursor宿主额度限制切换到Grok，由Grok重建和终检，按混合成品评价。第二轮无独立运行。')
lead('Kimi','第一轮多次提前结束，经人工续跑后退出；第二轮反复读取同一清单，评测方终止运行。两轮均未形成完整成品。')

h('2 两轮评分及主要差异',True)
p('成果质量Q与过程可靠性R各满分100分。Q包括数据判断、证据计算、调查推进、最终交付、复杂度与范围五个维度；R包括独立承担、自我纠错、路线控制、完成真实性、协议纪律五项。')
fig('Q1','表1  第一轮成果质量','分维度比较 A1:H10',16.8)
fig('R1','表2  第一轮过程可靠性','分维度比较 A25:H34',16.8)
lead('第一轮','Grok的质量与可靠性均居首，Q为94、R为99；GPT分别为93、97。Gemini为49、67，与其余五个独立交付系统差距较大。Claude的混合成品Q为81，Kimi未完成交付。')
h('2.1 第二轮评分',True)
fig('Q2','表3  第二轮成果质量','分维度比较 A13:H22',16.8)
fig('R2','表4  第二轮过程可靠性','分维度比较 A37:H46',16.8)
lead('成果质量','正式五系统的证据计算为21—25分，相差4分。含Gemini观察评分时，证据计算为16—25分、调查推进为23—30分、最终交付为12—15分。GPT质量100、Grok98，差异来自证据计算和最终交付。Qwen分项90，按主线时间窗错误封顶59；Gemini78仅作观察。')
lead('过程可靠性','Grok自我纠错得25分，GPT24，DeepSeek与GLM23，Qwen22，Gemini16。Gemini路线控制得满分，但仍有语义与解释错误。Claude不发布独立R，Kimi缺失分数留空。')

h('3 两轮得分与具体依据',True)
p('下列箭头依次表示第一轮、第二轮。Q为成果质量，R为过程可靠性。')
lead('GPT  Q 93→100  R 97→98','新增雪冰峰日、固定构成和替代解释检查，调查推进增加6分。峰日明细与排除对照、固定构成检查支持C2和C4满分，撤销原来缺少依据的两分扣分。计算与最终交付也为满分。')
lead('Grok  Q 94→98  R 99→99','质量增加的4分全部来自调查推进。地块排除比较和成熟队列支持深入调查；BBL地块被表述为单栋建筑，解释边界扣1分，恒真的文件存在检查使最终核验扣1分。R总分不变，自我纠错升1分、完成真实性降1分。')
lead('DeepSeek  Q 89→95  R 92→94','地址与固定观察期检查使调查推进达到30分。第二轮将最大创建时间写成02:22:25，核验值为01:50:33，自检未发现，最终文件检查仅得2/4。R增加来自路线控制改善。')
lead('GLM  Q 86→94  R 94→96','增量分解和热点调查使调查推进增加7分。月均值与关闭率修正提高过程评分，但最终说明仍将44个字段写成43个，最终交付由15分降至14分。Closed子集547条负时长已经核实，不作为错误扣分。')
lead('Qwen  Q 81→59  R 89→95','异常和渠道调查使调查推进增加7分，自我纠错增加4分。Use Indoor份额由约10.69%升至29.73%，与“结构无实质变化”的表述冲突；关闭时长主要发现使用11个月与9个月比较，分项合计90，按冻结规则封顶为59。较重的重建和核对投入使复杂度与范围减少2分。')
lead('Gemini  观察评分 Q 49→78  R 67→81','字段检查和贡献分解明显改善，但arrest关键词命中率8.28%包含否定等语境，明确肯定短语仅约0.013718%。将宽泛命中率解释为拘捕情况导致解释边界仅得1/5，交付后登记修正5项。第二轮越界搜索目录、模型子版本不可核，不进入正式配对汇总。')
lead('Claude  第一轮混合成品 Q 81','调查推进得25分，证据计算仅16分。地址聚集被认定为无效记录、部分因果解释缺证据，前20类增量与全量未对齐。Grok参与最终交付，故不评价Claude的独立过程可靠性。')
lead('Kimi  未评分','缺少完整成品与足够评分材料。第二轮真实复测14个实质动作中13个无收益，用时约2.65分钟。')

h('4 两轮变化与投入比较',True)
fig('分项增量','表5  五个配对系统的分项增量与封顶调整','评测总览 B29:N38',16.8)
p('五系统分项合计平均增加6.8分，调查推进贡献76.5%，最终交付合计减少1分。Qwen封顶调整−31分后，最终Q平均仅增加0.6分，中位增加6分。Gemini仅保留观察评分。')
fig('总览-第二轮','表6  第二轮交付与资源指标','评测总览 B8:N16 第二轮',16.8)
p('GPT未登记交付后修正，但可估算费用最高，为38.96元。GLM约9.6分钟、6.48元，Q与R均高于Qwen；DeepSeek费用最低，为1.37—2.74元，仍需重点核对终稿数字。Grok和Gemini缺少Token记录，费用不可比较。')
p('六个系统第二轮均在无追加指令下完成交付，后续仍登记10项修正。Qwen与Gemini的无收益动作占比均为0，仍分别有2项和5项修正，文件检查未充分覆盖解释错误。')
p('第二轮增加了调查和对照，但Qwen的比较月份、Gemini的文本解释等问题仍留在终稿。')

conclusions=json.loads((A/'评测结论.json').read_text(encoding='utf-8'))
h('5 评测结论与系统评价',True)
p(conclusions['总判断'])
for ai,judgement in conclusions['系统评价'].items():lead(ai,judgement)


h('6 模型与运行条件',True)
fig('运行条件','表7  模型、思考强度与Harness','字段说明 H1:L14',17)
p('海外模型思考强度设为中，国内模型设为Claude Code中的high。')
lead('Claude空白项原因','第一轮因Cursor宿主额度限制，由Claude Opus 5切换到Grok收尾，81分仅评价混合成品，独立过程R留空。第二轮没有可验证的独立运行，因此第二轮分数和两轮变化留空。')
lead('Kimi空白项原因','第一轮反复提前结束，需要人工续跑，随后退出；第二轮陷入重复读取同一清单，评测方确认循环后终止。两轮均无完整成品，无法评定Q与R。')

normalize_fonts(d)
d.save(A/'项目2_简洁版报告.docx')
(A/'简洁版图表来源.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('简洁版已生成，7个计划页面，7张Excel原表')
