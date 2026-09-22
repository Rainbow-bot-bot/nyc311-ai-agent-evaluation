# -*- coding: utf-8 -*-
"""生成 最终成果.xlsx 看板（读取 汇总/ 下 CSV，openpyxl 原生图表）"""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import LineChart, BarChart, PieChart, Reference

BASE = os.path.dirname(os.path.abspath(__file__))
H = os.path.join(BASE, "汇总")
load = lambda n: pd.read_csv(os.path.join(H, n))

wb = Workbook()

# ---- 样式 ----
TITLE_F = Font(size=16, bold=True, color="1F4E79")
H2_F = Font(size=12, bold=True, color="1F4E79")
HDR_F = Font(bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
NOTE_F = Font(size=9, color="808080", italic=True)
BODY_F = Font(size=10)
THIN = Border(*[Side(style="thin", color="D9D9D9")]*4)

def put_title(ws, cell, text):
    ws[cell] = text; ws[cell].font = TITLE_F

def put_h2(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text); c.font = H2_F
    return row + 1

def put_table(ws, df, start_row, start_col=1, num_fmt=None, max_rows=None):
    """写表头+数据，返回结束行"""
    if max_rows: df = df.head(max_rows)
    for j, colname in enumerate(df.columns):
        c = ws.cell(row=start_row, column=start_col+j, value=str(colname))
        c.font = HDR_F; c.fill = HDR_FILL; c.border = THIN
        c.alignment = Alignment(horizontal="center")
    for i, (_, r) in enumerate(df.iterrows()):
        for j, v in enumerate(r):
            if pd.isna(v): v = None
            elif hasattr(v, "item"): v = v.item()
            c = ws.cell(row=start_row+1+i, column=start_col+j, value=v)
            c.font = BODY_F; c.border = THIN
            if num_fmt and isinstance(v, (int, float)) and num_fmt.get(j):
                c.number_format = num_fmt[j]
    return start_row + len(df)

def put_notes(ws, row, lines, col=1):
    for t in lines:
        c = ws.cell(row=row, column=col, value=t); c.font = NOTE_F; row += 1
    return row

def autowidth(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

# ============ Sheet 1: 说明与概况 ============
ws = wb.active; ws.title = "说明与概况"
put_title(ws, "A1", "NYC 311 服务请求分析看板（2024-09 至 2026-09）")
ws["A2"] = "数据来源：D:\\项目1\\原始数据（25 个按月切分的 Parquet 分片，只读）｜生成日期：2026-09-16｜分析：Qwen"
ws["A2"].font = NOTE_F
r = 4
r = put_h2(ws, r, 1, "一、数据规模与范围")
ov = load("概况_总量.csv").T.reset_index()
ov.columns = ["项目", "值"]
r = put_table(ws, ov, r) + 1
ws.cell(row=r, column=1, value="说明：每行 = 一条 311 服务请求记录（unique_key 全表唯一，无重复）。首尾月份（2024-09、2026-09）为不完整月，同比等对比均使用完整月窗口。").font = NOTE_F
r += 2

r = put_h2(ws, r, 1, "二、关键字段缺失率")
r = put_table(ws, load("概况_缺失率.csv"), r) + 1
ws.cell(row=r, column=1, value="缺失影响：closed_date 缺失 3.19% 即未关闭工单（响应时长分析剔除）；due_date 缺失 99.6% 基本不可用；其余关键字段完整。").font = NOTE_F
r += 2

r = put_h2(ws, r, 1, "三、工单状态分布")
r = put_table(ws, load("概况_状态.csv"), r) + 2
r = put_h2(ws, r, 1, "四、受理渠道分布")
r = put_table(ws, load("概况_渠道.csv"), r) + 2
r = put_h2(ws, r, 1, "五、承办机构 TOP12")
r = put_table(ws, load("概况_机构.csv"), r)
autowidth(ws, {"A": 34, "B": 40, "C": 14, "D": 12})

# ============ Sheet 2: 核心发现 ============
ws = wb.create_sheet("核心发现")
put_title(ws, "A1", "五项核心发现（证据详见后续各表）")
findings = [
    ("发现1｜2025-26 年是明显的严冬：三类连锁反应同时出现",
     ["Snow or Ice（积雪/结冰）投诉：2024-12~2025-02 共 8,794 条 → 2025-12~2026-02 共 63,722 条，约 7.2 倍。",
      "HEAT/HOT WATER（供暖/热水）投诉：供暖季（10月~次年3月）256,140 → 312,448 条，+22.0%；布朗克斯区最多（10.8万条）。",
      "坑洞（Pothole）投诉：2026-03 单月 22,790 条，是 2025-03（4,313 条）的 5.3 倍；此后 4~8 月仍维持 4,300~11,400 条/月，高于往年基线（约1,700~4,300）。",
      "解读：三者时间与量级相互印证（融雪冻融→路面破损、低温→供暖故障），属于天气驱动的关联现象，非因果链的全部证据。"]),
    ("发现2｜违停投诉是全市第一大类，但约 2/3 到场后未采取执法行动",
     ["Illegal Parking（违章停车）两年共 1,159,053 条，占全部记录 15.4%，全部由 NYPD 承办。",
      "处置结构：到场未见违规 34.1%、开具传票 22.5%、判定无需警察行动 16.2%、到场时车辆已离开 13.5%、到场采取行动 11.4%。",
      "布区分布高度集中：布鲁克林 39.4% + 皇后区 31.2%，两区合计约 70.6%。",
      "注意：这里的百分比是『记录数』口径；同一辆车/同一地点可能产生多条记录。"]),
    ("发现3｜Drug Activity（涉毒活动）投诉自 2025-11 起骤降约七成，疑似受理口径变化",
     ["月量走势：2024-09~2025-06 基线约 1,500~2,500 条/月；2025-07~08 冲高至约 6,000 条/月；2025-11 起跌至 520~1,150 条/月。",
      "同比（2024-10~2025-08 vs 2025-10~2026-08）：27,869 → 8,955 条，-67.9%，为全部主要类型中降幅最大。",
      "细分描述仍为 Use Outside / Use Indoor，结构未变，故更像受理/派发流程调整而非现实事件消失；数据本身无法区分这两种解释。"]),
    ("发现4｜直升机噪音投诉自 2026-06 起骤降约八成",
     ["月量走势：2024-09~2026-05 稳定在约 1,000~2,850 条/月；2026-06 起降至 525、296、229 条/月（6、7、8月）。",
      "降幅集中在『Other』描述项；NYPD/News Gathering 等描述本就量小。",
      "推断（未验证）：可能与 2026 年年中直升机航线的管制/政策变化有关，数据内无法证实。"]),
    ("发现5｜响应时长呈两个世界：调度型约 1~2 小时，维修型以周计",
     ["NYPD/噪音/违停类（调度型）：中位 0.4~2.2 小时关闭——快是因为按『到场处置』即关闭，不代表问题已解决（见发现2）。",
      "HPD 房屋维修类：HEAT/HOT WATER 中位 41.8 小时，PLUMBING 190 小时，DOOR/WINDOW 288 小时，UNSANITARY CONDITION 250 小时。",
      "P90 尾部很长：PLUMBING、WATER LEAK 等维修类 P90 超过 1,300 小时（约 55 天）。",
      "口径：仅统计已关闭且 closed≥created 的记录（约 727 万条，负时长异常仅 1,488 条已剔除）。"]),
]
r = 3
for title, lines in findings:
    ws.cell(row=r, column=1, value=title).font = H2_F; r += 1
    for ln in lines:
        c = ws.cell(row=r, column=1, value="• " + ln); c.font = BODY_F
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 16 * (1 + len(ln)//60)
        r += 1
    r += 1
ws.cell(row=r, column=1, value="总体背景：两年 752.5 万条记录，月均约 31 万条；NYPD 承办 45.2%（噪音+违停为主），HPD 承办 21.7%（房屋维修类）。渠道以 ONLINE(44.3%)、PHONE(25.7%)、MOBILE(22.1%) 为主。").font = NOTE_F
autowidth(ws, {"A": 110})

# ============ Sheet 3: 月度趋势 ============
ws = wb.create_sheet("月度趋势")
put_title(ws, "A1", "月度记录量与关键类型走势")
m = load("月度_总量.csv")
r = put_h2(ws, 3, 1, "全部记录：月度总量")
end = put_table(ws, m, r)
ch = LineChart(); ch.title = "311 月度记录量（2024-09 ~ 2026-09）"; ch.height = 8; ch.width = 24
ch.add_data(Reference(ws, min_col=2, min_row=r, max_row=end), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r+1, max_row=end))
ch.y_axis.title = "记录数"; ws.add_chart(ch, f"K{r}")
r2 = end + 2

kt = load("月度_关键类型.csv")
ph = load("月度_坑洞.csv")
kt = kt.merge(ph, on="月份")
r3 = put_h2(ws, r2, 1, "关键投诉类型：月度量（含坑洞细分）")
end3 = put_table(ws, kt, r3)
ch2 = LineChart(); ch2.title = "天气相关类型走势（HEAT/HOT WATER、Snow or Ice、Street Condition、坑洞）"; ch2.height = 9; ch2.width = 26
for col in [2, 3, 4, 9]:  # HEAT, Snow, StreetCond, 坑洞
    ch2.add_data(Reference(ws, min_col=col, min_row=r3, max_row=end3), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=r3+1, max_row=end3))
ws.add_chart(ch2, f"K{r3}")
r4 = end3 + 2
ch3 = LineChart(); ch3.title = "结构性变化类型（Drug Activity、Noise - Helicopter）"; ch3.height = 8; ch3.width = 26
for col in [7, 8]:
    ch3.add_data(Reference(ws, min_col=col, min_row=r3, max_row=end3), titles_from_data=True)
ch3.set_categories(Reference(ws, min_col=1, min_row=r3+1, max_row=end3))
ws.add_chart(ch3, f"K{r4+18}")
put_notes(ws, end3+18, ["注：2024-09 自 9-07 起、2026-09 至 9-05 止，均为不完整月。",
                        "严冬信号：Snow or Ice 2026-01/02 达 2.4万/3.1万条（上年同期 0.5万/0.3万）；坑洞 2026-03 达 2.28万条为全程峰值。"])
autowidth(ws, {"A": 12, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16, "H": 18, "I": 12})

# ============ Sheet 4: 类型TOP与同比 ============
ws = wb.create_sheet("类型TOP与同比")
put_title(ws, "A1", "投诉类型 TOP20 与同比变化")
t = load("类型_TOP20.csv")
r = put_h2(ws, 3, 1, "记录量 TOP20 投诉类型（全期）")
end = put_table(ws, t, r, num_fmt={2: "#,##0"})
ch = BarChart(); ch.type = "bar"; ch.title = "TOP20 投诉类型（记录数）"; ch.height = 14; ch.width = 22
ch.add_data(Reference(ws, min_col=3, min_row=r, max_row=end), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r+1, max_row=end))
ch.legend = None; ws.add_chart(ch, f"F{r}")
r2 = end + 2

y = load("同比_类型变化.csv")
r3 = put_h2(ws, r2, 1, "同比变化 TOP15（上升）｜口径：2024-10~2025-08 vs 2025-10~2026-08，各11个完整月，仅含两年合计>2万的类型")
end3 = put_table(ws, y.head(15), r3, num_fmt={1: "#,##0", 2: "#,##0", 3: "#,##0"})
r4 = end3 + 2
r4 = put_h2(ws, r4, 1, "同比变化 TOP15（下降）")
end4 = put_table(ws, y.tail(15).sort_values("变化率pct"), r4, num_fmt={1: "#,##0", 2: "#,##0", 3: "#,##0"})
ch2 = BarChart(); ch2.type = "bar"; ch2.title = "同比变化率%（升TOP15 / 降TOP15）"; ch2.height = 16; ch2.width = 22
ch2.add_data(Reference(ws, min_col=5, min_row=r3, max_row=end3), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=r3+1, max_row=end3))
ws.add_chart(ch2, f"G{r3}")
ch3 = BarChart(); ch3.type = "bar"
ch3.add_data(Reference(ws, min_col=5, min_row=r4, max_row=end4), titles_from_data=True)
ch3.set_categories(Reference(ws, min_col=1, min_row=r4+1, max_row=end4))
ch3.legend = None; ws.add_chart(ch3, f"G{r4}")
autowidth(ws, {"A": 32, "B": 20, "C": 14, "D": 12, "E": 12})

# ============ Sheet 5: 行政区 ============
ws = wb.create_sheet("行政区")
put_title(ws, "A1", "行政区分布")
b = load("行政区_总量.csv")
r = put_h2(ws, 3, 1, "各行政区记录总量")
end = put_table(ws, b, r, num_fmt={1: "#,##0"})
ch = BarChart(); ch.title = "各行政区记录数"; ch.height = 8; ch.width = 16
ch.add_data(Reference(ws, min_col=2, min_row=r, max_row=end), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r+1, max_row=end))
ch.legend = None; ws.add_chart(ch, f"I{r}")
r2 = end + 2

bx = load("行政区_类型交叉.csv")
cross = bx.pivot(index="投诉类型", columns="行政区", values="记录数").fillna(0).astype(int).reset_index()
r3 = put_h2(ws, r2, 1, "行政区 × 六大类型交叉（记录数）")
end3 = put_table(ws, cross, r3, num_fmt={i: "#,##0" for i in range(1, len(cross.columns))})
r4 = end3 + 2

vp = load("违停_行政区.csv")
r5 = put_h2(ws, r4, 1, "违停投诉的行政区分布（占违停总数比例）")
end5 = put_table(ws, vp, r5)
ch2 = PieChart(); ch2.title = "违停投诉行政区占比"; ch2.height = 8; ch2.width = 12
ch2.add_data(Reference(ws, min_col=2, min_row=r5, max_row=end5), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=r5+1, max_row=end5))
ws.add_chart(ch2, f"E{r5}")
put_notes(ws, end5+2, ["注：此处为记录数口径，未按人口/面积标准化；布鲁克林、皇后区本身人口与街区密度高，占比高不必然代表『更严重』。"])
autowidth(ws, {"A": 26, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16})

# ============ Sheet 6: 响应与处置 ============
ws = wb.create_sheet("响应与处置")
put_title(ws, "A1", "响应时长与处置结构")
rt = load("响应_类型.csv")
r = put_h2(ws, 3, 1, "各类型关闭时长（已关闭记录，created→closed，小时）")
end = put_table(ws, rt, r, num_fmt={1: "#,##0"})
ch = BarChart(); ch.type = "bar"; ch.title = "中位关闭时长（小时，对数量级差异大请结合P90看）"; ch.height = 16; ch.width = 20
ch.add_data(Reference(ws, min_col=3, min_row=r, max_row=end), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r+1, max_row=end))
ch.legend = None; ws.add_chart(ch, f"G{r}")
r2 = end + 2

vp = load("违停_处置结构.csv")
r3 = put_h2(ws, r2, 1, "违停投诉处置结构（按 resolution_description 关键词归类）")
end3 = put_table(ws, vp, r3)
ch2 = PieChart(); ch2.title = "违停处置结构"; ch2.height = 9; ch2.width = 13
ch2.add_data(Reference(ws, min_col=2, min_row=r3, max_row=end3), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=r3+1, max_row=end3))
ws.add_chart(ch2, f"E{r3}")
r4 = end3 + 2

hs = load("供暖_两季对比.csv")
hs["变化率%"] = (100*(hs["供暖季2025_26"]/hs["供暖季2024_25"]-1)).round(1)
r5 = put_h2(ws, r4, 1, "供暖投诉两季对比（HEAT/HOT WATER，10月~次年3月）")
end5 = put_table(ws, hs, r5, num_fmt={1: "#,##0", 2: "#,##0"})
ch3 = BarChart(); ch3.title = "供暖投诉：2024-25季 vs 2025-26季"; ch3.height = 8; ch3.width = 16
ch3.add_data(Reference(ws, min_col=2, max_col=3, min_row=r5, max_row=end5), titles_from_data=True)
ch3.set_categories(Reference(ws, min_col=1, min_row=r5+1, max_row=end5))
ws.add_chart(ch3, f"F{r5}")
r6 = end5 + 2
wt = load("冬季_两季对比.csv")
wt["倍数"] = (wt["冬2025_26"]/wt["冬2024_25"]).round(1)
r7 = put_h2(ws, r6, 1, "冬季除雪类对比（Snow or Ice，12月~次年2月）")
end7 = put_table(ws, wt, r7, num_fmt={1: "#,##0", 2: "#,##0"})
put_notes(ws, end7+2, ["注：关闭时长=系统关闭工单的时间差，衡量『处理流程速度』而非『问题彻底解决』；调度型工单到场即关，故中位数很小。"])
autowidth(ws, {"A": 32, "B": 16, "C": 16, "D": 14, "E": 14, "F": 14, "G": 14})

for s in wb.worksheets:
    s.freeze_panes = "A2"

path = os.path.join(BASE, "最终成果.xlsx")
wb.save(path)
print("saved:", path)
