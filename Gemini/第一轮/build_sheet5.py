import pandas as pd
from openpyxl.chart import LineChart, Reference
from common_styles import *

def build_sheet_5(wb):
    print("Building Sheet 5: 发现3_违停与噪音时空节律...")
    hourly_df = pd.read_parquet("D:/项目2/Gemini/data_cache/hourly.parquet")
    dow_df = pd.read_parquet("D:/项目2/Gemini/data_cache/dow.parquet")
    resolution_actions_df = pd.read_parquet("D:/项目2/Gemini/data_cache/resolution_actions.parquet")
    
    ws5 = wb.create_sheet(title="发现3_违停与噪音时空节律")
    ws5.row_dimensions[1].height = 28
    ws5.row_dimensions[2].height = 18

    ws5.cell(1, 1, "核心发现 3：城市摩擦的双面性 —— 违章停车与各类噪音的昼夜与周末交错节律").font = font_title
    ws5.cell(2, 1, "核心结论：噪音与违停合计占全市 41% 诉求；噪音呈‘夜间与周末派对爆发型’，违停呈‘早晚通勤潮汐型’，交替挤占城市执法资源").font = font_subtitle

    # Table 5-1: 24-Hour Diurnal Rhythm
    ws5.cell(4, 1, "【表 5-1 城市民生诉求 24 小时昼夜交替节律全景表】").font = font_sec_head
    ws5.row_dimensions[5].height = 22
    headers_hr = ["时段 (整点)", "全市民生工单量", "各类噪音工单", "违停阻道工单", "供暖热水工单", "噪音占该时段比重", "违停占该时段比重", "主导城市运行摩擦特征"]
    for i, h in enumerate(headers_hr, 1):
        c = ws5.cell(5, i, h)
        c.font = font_th
        c.fill = fill_th
        c.alignment = align_center
        c.border = border_cell

    for idx, r in hourly_df.iterrows():
        row_num = 6 + idx
        ws5.row_dimensions[row_num].height = 19
        hr = int(r['hour_of_day'])
        tot = r['total_requests']
        n_cnt = r['noise_requests']
        p_cnt = r['parking_requests']
        h_cnt = r['heat_requests']
        n_p = n_cnt / tot
        p_p = p_cnt / tot
        
        if hr in (22, 23, 0):
            feat = "【夜间噪音极值爆发期】居民音乐/派对扰民达顶峰，噪音占比接近或超过 50%！"
        elif hr in (1, 2, 3, 4):
            feat = "【深夜安静期】全城诉求总量降至全天谷底，偶发夜间酒吧与街头噪音"
        elif hr in (7, 8, 9):
            feat = "【早通勤违停高峰期】上班送学车流启动，堵塞私家车道、占消防栓投诉剧增"
        elif hr in (10, 11, 12, 13, 14):
            feat = "【日间常规运行期】市政综合巡查与日常维修诉求平稳释放"
        elif hr in (18, 19, 20):
            feat = "【晚通勤归家潮】下班抢车位纠纷上升，夜生活社交活动开启"
        else:
            feat = "【过渡转换期】诉求结构平稳轮转"
            
        vals = [
            f"{hr:02d}:00 - {hr:02d}:59",
            tot,
            n_cnt,
            p_cnt,
            h_cnt,
            n_p,
            p_p,
            feat
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws5.cell(row_num, c_idx, val)
            cell.font = font_td
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_center
            elif c_idx in (2, 3, 4, 5):
                cell.alignment = align_right
                cell.number_format = "#,##0"
            elif c_idx in (6, 7):
                cell.alignment = align_right
                cell.number_format = "0.0%"
                if c_idx == 6 and val > 0.40:
                    cell.fill = fill_alert
                    cell.font = font_td_bold
                elif c_idx == 7 and val > 0.20:
                    cell.fill = fill_highlight
                    cell.font = font_td_bold
            else:
                cell.alignment = align_left
            if row_num % 2 == 1 and cell.fill not in (fill_alert, fill_highlight):
                cell.fill = fill_zebra

    # Total row for hourly
    tot_r_hr = 6 + len(hourly_df)
    ws5.cell(tot_r_hr, 1, "全天 24 小时汇总").font = font_total
    ws5.cell(tot_r_hr, 1).alignment = align_center
    ws5.cell(tot_r_hr, 1).fill = fill_total
    ws5.cell(tot_r_hr, 1).border = border_total

    for c_idx, col_let in enumerate(['B', 'C', 'D', 'E'], 2):
        cell = ws5.cell(tot_r_hr, c_idx, f"=SUM({col_let}6:{col_let}{tot_r_hr-1})")
        cell.font = font_total
        cell.alignment = align_right
        cell.fill = fill_total
        cell.border = border_total
        cell.number_format = "#,##0"

    cell_np = ws5.cell(tot_r_hr, 6, f"=C{tot_r_hr}/B{tot_r_hr}")
    cell_np.font = font_total
    cell_np.alignment = align_right
    cell_np.fill = fill_total
    cell_np.border = border_total
    cell_np.number_format = "0.0%"

    cell_pp = ws5.cell(tot_r_hr, 7, f"=D{tot_r_hr}/B{tot_r_hr}")
    cell_pp.font = font_total
    cell_pp.alignment = align_right
    cell_pp.fill = fill_total
    cell_pp.border = border_total
    cell_pp.number_format = "0.0%"

    ws5.cell(tot_r_hr, 8, "全天昼夜互补：白昼违停争路权，夜晚噪音扰清梦").fill = fill_total
    ws5.cell(tot_r_hr, 8).border = border_total
    ws5.cell(tot_r_hr, 8).font = font_total

    # Table 5-2: Day of Week Rhythm
    start_r_dow = 33
    ws5.cell(start_r_dow, 1, "【表 5-2 星期周中与周末生活节律对比表】").font = font_sec_head
    ws5.row_dimensions[start_r_dow+1].height = 22
    headers_dow = ["星期 (Day of Week)", "全市民生工单量", "各类噪音工单", "违停阻道工单", "供暖热水工单", "噪音占当日比重", "较周中平均的周末增幅", "周中周末运行特征研判"]
    for i, h in enumerate(headers_dow, 1):
        c = ws5.cell(start_r_dow+1, i, h)
        c.font = font_th
        c.fill = fill_th_sec
        c.alignment = align_center
        c.border = border_cell

    dow_order = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
    zh_dow = {
        'Monday': '星期一 (Monday)', 'Tuesday': '星期二 (Tuesday)', 'Wednesday': '星期三 (Wednesday)',
        'Thursday': '星期四 (Thursday)', 'Friday': '星期五 (Friday)', 'Saturday': '星期六 (Saturday)', 'Sunday': '星期日 (Sunday)'
    }

    weekday_noise_avg = 193906.0
    for idx, d_name in enumerate(dow_order):
        row_num = start_r_dow + 2 + idx
        ws5.row_dimensions[row_num].height = 20
        sub_r = dow_df[dow_df['day_of_week'] == d_name].iloc[0]
        tot = sub_r['total_requests']
        n_cnt = sub_r['noise_requests']
        p_cnt = sub_r['parking_requests']
        h_cnt = sub_r['heat_requests']
        n_p = n_cnt / tot
        
        diff_pct = (n_cnt - weekday_noise_avg) / weekday_noise_avg
        if d_name in ('Saturday', 'Sunday'):
            eval_w = f"【周末狂欢潮】较周中激增 {diff_pct*100:.1f}%，夜间派对与街头聚会密集爆发"
        elif d_name == 'Friday':
            eval_w = "【周末前哨战】周五傍晚起噪音投诉开始抬头"
        else:
            eval_w = "【工作日平稳期】以通勤违停、市容环卫及房屋报修为主流"
            
        vals = [
            zh_dow[d_name],
            tot,
            n_cnt,
            p_cnt,
            h_cnt,
            n_p,
            f"+{diff_pct*100:.1f}%" if diff_pct > 0 else f"{diff_pct*100:.1f}%",
            eval_w
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws5.cell(row_num, c_idx, val)
            cell.font = font_td
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_left
                if d_name in ('Saturday', 'Sunday'):
                    cell.fill = fill_alert
                    cell.font = font_td_bold
            elif c_idx in (2, 3, 4, 5):
                cell.alignment = align_right
                cell.number_format = "#,##0"
                if d_name in ('Saturday', 'Sunday') and c_idx == 3:
                    cell.fill = fill_alert
                    cell.font = font_td_bold
            elif c_idx == 6:
                cell.alignment = align_right
                cell.number_format = "0.0%"
                if d_name in ('Saturday', 'Sunday'):
                    cell.fill = fill_alert
                    cell.font = font_td_bold
            elif c_idx == 7:
                cell.alignment = align_center
                if d_name in ('Saturday', 'Sunday'):
                    cell.fill = fill_alert
                    cell.font = font_td_bold
            else:
                cell.alignment = align_left
            if row_num % 2 == 1 and cell.fill != fill_alert:
                cell.fill = fill_zebra

    # Table 5-3: NYPD Resolution Breakdown for Noise & Parking
    start_r_res = 44
    ws5.cell(start_r_res, 1, "【表 5-3 警务局 (NYPD) 现场处置结果深度解构：违停与噪音的执法实质】").font = font_sec_head
    ws5.row_dimensions[start_r_res+1].height = 22
    headers_res = ["诉求主类别", "现场处置官方认定结果大类", "处置案件量 (件)", "类内占比 (%)", "现场执法实质评价与空转研判"]
    for i, h in enumerate(headers_res, 1):
        c = ws5.cell(start_r_res+1, i, h)
        c.font = font_th
        c.fill = fill_th
        c.alignment = align_center
        c.border = border_cell

    action_evals = {
        ('Noise - Residential', '未发现违规证据 (No Evidence Observed)'): "【典型空转】警察到场时噪音已停息，或声响未达法定噪音分贝标准 (占 44.9%)",
        ('Noise - Residential', '其他处理结果 (Other Resolutions)'): "包括转交辖区社区事务组(Community Affairs)或无管辖权解释",
        ('Noise - Residential', '已采取处置措施修复 (Action Taken to Fix)'): "【有效干预】口头告诫劝导、勒令降低音量或劝散聚会人员 (占 17.3%)",
        ('Noise - Residential', '判定无需警察执法 (Action Not Necessary)'): "判定属于正常生活起居声响（如脚步声、儿童玩闹），无需公权力介入",
        ('Noise - Residential', '到达现场时当事人已离开 (Condition/Party Gone)'): "多见于楼下流动喧哗，巡警到场前已自行散去",
        ('Noise - Residential', '开具传票/罚单 (Summons Issued)'): "【极低概率】仅 0.36% 的恶劣屡教不改案例才会由警方开出刑事传票或民事罚单！",
        ('Illegal Parking', '未发现违规证据 (No Evidence Observed)'): "车主已挪车、车辆停放合规或报送地址有误 (占 25.9%)",
        ('Illegal Parking', '开具传票/罚单 (Summons Issued)'): "【实质执法】警方确认违章并现场张贴停车罚单 (占 22.5%)",
        ('Illegal Parking', '判定无需警察执法 (Action Not Necessary)'): "不属于即时拖车或违规范畴，判定不予处置",
        ('Illegal Parking', '到达现场时当事人已离开 (Condition/Party Gone)'): "【时间差空转】违停车辆在巡警到场前已驶离，扑空率 13.5%",
        ('Illegal Parking', '已采取处置措施修复 (Action Taken to Fix)'): "口头警告车主立即驶离，消除了道路阻滞",
        ('Illegal Parking', '其他处理结果 (Other Resolutions)'): "转交交通局拖车管理处或因天气暂缓执法"
    }

    for idx, r in resolution_actions_df.iterrows():
        row_num = start_r_res + 2 + idx
        ws5.row_dimensions[row_num].height = 20
        ctype = r['complaint_type']
        act = r['action_category']
        eval_txt = action_evals.get((ctype, act), "常规办结流程")
        vals = [
            "住宅噪音 (Noise - Residential)" if ctype == 'Noise - Residential' else "违章停车 (Illegal Parking)",
            act,
            r['count_cases'],
            r['pct_within_type'] / 100.0,
            eval_txt
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws5.cell(row_num, c_idx, val)
            cell.font = font_td
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_center
                if idx in (0, 6):
                    cell.fill = fill_highlight
            elif c_idx == 2:
                cell.alignment = align_left
            elif c_idx == 3:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            elif c_idx == 4:
                cell.alignment = align_right
                cell.number_format = "0.00%"
                if idx == 11 and ctype == 'Noise - Residential':
                    cell.fill = fill_alert
                    cell.font = font_td_bold
            else:
                cell.alignment = align_left
            if row_num % 2 == 1 and cell.fill not in (fill_alert, fill_highlight):
                cell.fill = fill_zebra

    # Embedded Diurnal Line Chart (Noise vs Parking)
    chart_diurnal = LineChart()
    chart_diurnal.title = "城市运行 24 小时昼夜双峰交错对比 (噪音 vs 违停)"
    chart_diurnal.style = 11
    chart_diurnal.y_axis.title = "工单数量 (件)"
    chart_diurnal.x_axis.title = "时段 (小时)"
    chart_diurnal.height = 13
    chart_diurnal.width = 19

    data_ref_d = Reference(ws5, min_col=3, min_row=5, max_col=4, max_row=29)
    cats_ref_d = Reference(ws5, min_col=1, min_row=6, max_row=29)
    chart_diurnal.add_data(data_ref_d, titles_from_data=True)
    chart_diurnal.set_categories(cats_ref_d)
    ws5.add_chart(chart_diurnal, "H5")

    autofit(ws5, 1, 8)
    ws5.column_dimensions['A'].width = 16
    ws5.column_dimensions['B'].width = 16
    ws5.column_dimensions['C'].width = 16
    ws5.column_dimensions['D'].width = 16
    ws5.column_dimensions['E'].width = 16
    ws5.column_dimensions['F'].width = 16
    ws5.column_dimensions['G'].width = 16
    ws5.column_dimensions['H'].width = 46
