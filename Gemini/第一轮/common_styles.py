import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Segoe UI"
font_title = Font(name=FONT_NAME, size=15, bold=True, color="1F4E79")
font_subtitle = Font(name=FONT_NAME, size=9, italic=True, color="595959")
font_sec_head = Font(name=FONT_NAME, size=11, bold=True, color="1F4E79")
font_th = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
font_td = Font(name=FONT_NAME, size=9, color="1F2937")
font_td_bold = Font(name=FONT_NAME, size=9, bold=True, color="1F2937")
font_total = Font(name=FONT_NAME, size=9, bold=True, color="1F4E79")

fill_th = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
fill_th_sec = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_total = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
fill_card = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
fill_alert = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
fill_success = PatternFill(start_color="DEF7EC", end_color="DEF7EC", fill_type="solid")

thin_border_side = Side(style="thin", color="D1D5DB")
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_total = Border(
    left=thin_border_side, right=thin_border_side,
    top=Side(style="thin", color="1F4E79"),
    bottom=Side(style="double", color="1F4E79")
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

def autofit(ws, min_col=1, max_col=None, max_len_cap=60):
    if max_col is None:
        max_col = ws.max_column
    for col_idx in range(min_col, max_col + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            val = ws.cell(row_idx, col_idx).value
            if val is not None:
                s = str(val)
                line_len = sum(2 if ord(c) > 127 else 1 for c in s)
                if line_len > max_len:
                    max_len = line_len
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), max_len_cap)
    ws.views.sheetView[0].showGridLines = True

def add_kpi(ws, c1, r1, c2, r2, title, val, sub, color="1F4E79"):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.fill = fill_card
            cell.border = Border(
                left=Side(style="medium", color=color) if c == c1 else thin_border_side,
                right=Side(style="thin", color="CBD5E1") if c == c2 else thin_border_side,
                top=Side(style="thin", color="CBD5E1") if r == r1 else thin_border_side,
                bottom=Side(style="thin", color="CBD5E1") if r == r2 else thin_border_side
            )
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r1, end_column=c2)
    ws.merge_cells(start_row=r1+1, start_column=c1, end_row=r1+1, end_column=c2)
    ws.merge_cells(start_row=r1+2, start_column=c1, end_row=r1+2, end_column=c2)
    
    t_cell = ws.cell(r1, c1)
    t_cell.value = title
    t_cell.font = Font(name=FONT_NAME, size=8, bold=True, color="64748B")
    t_cell.alignment = align_center
    
    v_cell = ws.cell(r1+1, c1)
    v_cell.value = val
    v_cell.font = Font(name=FONT_NAME, size=15, bold=True, color=color)
    v_cell.alignment = align_center
    
    s_cell = ws.cell(r1+2, c1)
    s_cell.value = sub
    s_cell.font = Font(name=FONT_NAME, size=7, italic=True, color="94A3B8")
    s_cell.alignment = align_center
